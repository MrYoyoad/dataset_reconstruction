"""Rebuild a .pptx from a spec produced by import_pptx.py (round-trip regeneration).

    python scripts/deck/build_from_spec.py <spec_dir> <out.pptx> [--fix-page-numbers]

Everything the importer modelled is rebuilt natively (text boxes with run-level formatting, auto-shapes with
fill/line/adjustments, pictures with crops, tables, backgrounds, speaker notes); connectors are re-inserted
from their captured XML. --fix-page-numbers rewrites a page-number box to the true slide index and deck length, but ONLY when the
denominator already equals the deck length and the numerator is <= it. That guard exists because the naive
version (any "N / M" box) destroyed real content: it overwrote two job-id pairs ("392821 / 390026",
"229722 / 237301") on an appendix slide. Note also that a deck may use LOGICAL page numbers (continuation
slides sharing one number, as v20 does with "n / 35" over 37 physical slides) — in that case the guard
correctly declines to renumber, and it must, because the table of contents references the logical numbers.
"""
import argparse
import copy
import json
import os
import re

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

_NUM = re.compile(r"\((\d+)\)")


def _enum(enum_cls, s, default=None):
    if not s:
        return default
    m = _NUM.search(s)
    if m:
        try:
            return enum_cls(int(m.group(1)))
        except Exception:
            pass
    name = s.split(".")[-1].split(" ")[0]
    try:
        return getattr(enum_cls, name)
    except Exception:
        return default


def _set_color(color_obj, spec):
    if not spec:
        return False
    if "rgb" in spec and spec["rgb"]:
        color_obj.rgb = RGBColor.from_string(spec["rgb"])
        return True
    return False


_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _clear_defaults(tf, spec):
    """python-pptx injects defaults the original may not have: add_textbox adds <a:spAutoFit/>, add_shape sets
    algn='ctr' and bodyPr anchor='ctr'. Where the spec says the attribute was unset, remove it so the rebuild
    matches the original's inherit-state instead of silently re-styling (audit F2/F3/F4)."""
    bodyPr = tf._txBody.find(_A + "bodyPr")
    if bodyPr is not None:
        if spec.get("auto_size") is None:
            for tag in ("spAutoFit", "normAutofit"):
                el = bodyPr.find(_A + tag)
                if el is not None:
                    bodyPr.remove(el)
        if spec.get("vertical_anchor") is None and "anchor" in bodyPr.attrib:
            del bodyPr.attrib["anchor"]
    for i, p in enumerate(spec.get("paragraphs") or []):
        if p.get("alignment") is None and i < len(tf.paragraphs):
            pPr = tf.paragraphs[i]._p.find(_A + "pPr")
            if pPr is not None and "algn" in pPr.attrib:
                del pPr.attrib["algn"]


def _apply_text(tf, spec):
    tf.word_wrap = spec.get("word_wrap")
    va = _enum(MSO_ANCHOR, spec.get("vertical_anchor"))
    if va is not None:
        tf.vertical_anchor = va
    m = spec.get("margins") or []
    if len(m) == 4:
        tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom = [Emu(x) if x is not None else Emu(0) for x in m]
    paras = spec.get("paragraphs") or []
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        al = _enum(PP_ALIGN, p.get("alignment"))
        if al is not None:
            para.alignment = al
        if p.get("level"):
            para.level = p["level"]
        if p.get("line_spacing") is not None:
            para.line_spacing = p["line_spacing"]
        if p.get("space_before") is not None:
            para.space_before = Pt(p["space_before"])
        if p.get("space_after") is not None:
            para.space_after = Pt(p["space_after"])
        for r in p.get("runs") or []:
            run = para.add_run()
            run.text = r.get("text", "")
            f = run.font
            if r.get("font"):
                f.name = r["font"]
            if r.get("size"):
                f.size = Pt(r["size"])
            for k in ("bold", "italic", "underline"):
                if r.get(k) is not None:
                    setattr(f, k, r[k])
            _set_color(f.color, r.get("color"))
    _clear_defaults(tf, spec)


def _add(slide, sh, media_dir):
    k = sh["kind"]
    L, T, W, H = (Emu(sh[x]) if sh.get(x) is not None else None for x in ("left", "top", "width", "height"))
    if k == "group":
        raise ValueError("GROUP shape: child geometry is in group-relative coordinates (chOff/chExt); flattening "
                         "would mis-place it. Ungroup in PowerPoint and re-import, or extend the builder.")
    if k == "picture":
        pic = slide.shapes.add_picture(os.path.join(media_dir, sh["media"]), L, T, W, H)
        c = sh.get("crop")
        if c and any(c):
            pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom = c
        return
    if k == "table":
        rows, cols = len(sh["rows"]), len(sh["rows"][0])
        gt = slide.shapes.add_table(rows, cols, L, T, W, H)
        tbl = gt.table
        tbl.first_row = sh.get("first_row", False)
        tbl.horz_banding = sh.get("horz_banding", False)
        for j, w in enumerate(sh.get("col_widths") or []):
            if w:
                tbl.columns[j].width = Emu(w)
        for i, h in enumerate(sh.get("row_heights") or []):
            if h:
                tbl.rows[i].height = Emu(h)
        for i, row in enumerate(sh["rows"]):
            for j, cell_spec in enumerate(row):
                cell = tbl.cell(i, j)
                cell.text_frame.clear()
                _apply_text(cell.text_frame, cell_spec)
                fc = (sh.get("cell_fills") or [[]])[i][j] if sh.get("cell_fills") else None
                if fc:
                    cell.fill.solid()
                    _set_color(cell.fill.fore_color, fc)
        return
    if k == "unsupported":
        el = etree.fromstring(sh["xml"])
        xml = sh["xml"]
        if 'r:embed' in xml or 'r:link' in xml or ('r:id' in xml and 'hlinkClick' not in xml):
            raise ValueError("shape carries a relationship (r:embed/r:id/r:link) — charts, media and OLE cannot be "
                             "rebuilt without copying their parts; refusing rather than emitting a broken file")
        used = {int(e.get("id")) for e in slide.shapes._spTree.iter() if e.tag.endswith("}cNvPr") and e.get("id", "").isdigit()}
        for e in el.iter():
            if e.tag.endswith("}cNvPr") and e.get("id", "").isdigit():
                if int(e.get("id")) in used:
                    e.set("id", str(max(used) + 1))
                used.add(int(e.get("id")))
        slide.shapes._spTree.append(el)
        return
    # textbox / auto-shape
    if k == "shape" and sh.get("autoshape") and "TEXT_BOX" not in sh["autoshape"]:
        st = _enum(MSO_SHAPE, sh.get("auto_shape_type") or sh["autoshape"], MSO_SHAPE.RECTANGLE)
        shp = slide.shapes.add_shape(st, L, T, W, H)
        shp.shadow.inherit = False
        fill = sh.get("fill") or {}
        if fill.get("type") == "solid" and fill.get("color"):
            shp.fill.solid()
            _set_color(shp.fill.fore_color, fill["color"])
        else:
            shp.fill.background()
        ln = sh.get("line") or {}
        if ln.get("color") and ln["color"].get("rgb"):
            _set_color(shp.line.color, ln["color"])
            if ln.get("width"):
                shp.line.width = Emu(ln["width"])
        else:
            shp.line.fill.background()
        adj = sh.get("adjustments") or []
        for i, a in enumerate(adj):
            try:
                shp.adjustments[i] = a
            except Exception:
                pass
    else:
        shp = slide.shapes.add_textbox(L, T, W, H)
    if sh.get("rotation"):
        shp.rotation = sh["rotation"]
    if sh.get("text"):
        _apply_text(shp.text_frame, sh["text"])


def build(spec_dir, out, fix_page_numbers=False):
    spec = json.load(open(os.path.join(spec_dir, "deck_spec.json")))
    media = os.path.join(spec_dir, "media")
    prs = Presentation()
    prs.slide_width, prs.slide_height = spec["slide_width"], spec["slide_height"]
    blank = prs.slide_layouts[6]
    total = len(spec["slides"])
    fixed = 0
    for idx, sl in enumerate(spec["slides"], 1):
        s = prs.slides.add_slide(blank)
        if sl.get("background") and sl["background"].get("rgb"):
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = RGBColor.from_string(sl["background"]["rgb"])
        for sh in sl["shapes"]:
            if fix_page_numbers and sh.get("kind") in ("textbox", "shape") and sh.get("text"):
                paras = sh["text"].get("paragraphs") or []
                flat = "".join(r.get("text", "") for p in paras for r in (p.get("runs") or []))
                mnum = re.fullmatch(r"\s*(\d+)\s*/\s*(\d+)\s*", flat)
                is_page_number = bool(mnum) and int(mnum.group(2)) == total and int(mnum.group(1)) <= total
                if is_page_number and paras and paras[0].get("runs"):
                    sh = copy.deepcopy(sh)
                    runs = sh["text"]["paragraphs"][0]["runs"]
                    runs[0]["text"] = f"{idx} / {total}"
                    for r in runs[1:]:
                        r["text"] = ""
                    fixed += 1
            _add(s, sh, media)
        if sl.get("notes"):
            s.notes_slide.notes_text_frame.text = sl["notes"]
    prs.save(out)
    print(f"[build] {total} slides -> {out}" + (f"  (page numbers fixed on {fixed} slides)" if fix_page_numbers else ""))
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("spec_dir")
    ap.add_argument("out")
    ap.add_argument("--fix-page-numbers", action="store_true")
    a = ap.parse_args()
    build(a.spec_dir, a.out, a.fix_page_numbers)
