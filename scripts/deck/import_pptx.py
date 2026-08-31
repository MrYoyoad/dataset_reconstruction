"""Import a hand-edited .pptx into a regenerable spec (JSON + extracted media).

Why this exists: the 2026-08-31 deck was finished by hand (v20, 37 slides) while the tracked generator
(build_deck_2026_08_31.py) still produced the 29-slide version — the file and the code diverged, and only
the file had the final content. This importer closes that loop: any pptx can be turned into a spec that
build_from_spec.py reproduces, so the deck stays reproducible even after manual editing.

    python scripts/deck/import_pptx.py <deck.pptx> <spec_dir>

Captures per slide: background fill, and for every shape (recursively through groups) its type, geometry,
z-order, fill/line, text (paragraph/run level: font, size, bold/italic, colour, alignment), table cells,
picture bytes (written to <spec_dir>/media/), plus the speaker notes. Anything it cannot model is recorded
with kind="unsupported" and its raw XML, so the builder can re-insert it verbatim.
"""
import argparse
import hashlib
import json
import os

from pptx import Presentation
from pptx.util import Emu

MSO_PICTURE, MSO_TABLE, MSO_GROUP = 13, 19, 6


def _color(c):
    try:
        if c is None or c.type is None:
            return None
        if str(c.type).startswith("MSO_THEME_COLOR"):
            return {"theme": str(c.theme_color), "brightness": c.brightness}
        return {"rgb": str(c.rgb)}
    except Exception:
        return None


def _fill(shape):
    try:
        f = shape.fill
        t = str(f.type)
        if "BACKGROUND" in t or f.type is None:
            return {"type": "none"}
        if "SOLID" in t:
            return {"type": "solid", "color": _color(f.fore_color)}
        return {"type": t}
    except Exception:
        return None


def _line(shape):
    try:
        ln = shape.line
        return {"color": _color(ln.color), "width": ln.width.emu if ln.width is not None else None,
                "fill_type": str(ln.fill.type) if ln.fill is not None else None}
    except Exception:
        return None


def _text(tf):
    paras = []
    for p in tf.paragraphs:
        runs = [{"text": r.text, "font": r.font.name, "size": r.font.size.pt if r.font.size else None,
                 "bold": r.font.bold, "italic": r.font.italic, "underline": r.font.underline,
                 "color": _color(r.font.color)} for r in p.runs]
        paras.append({"runs": runs, "alignment": str(p.alignment) if p.alignment else None,
                      "level": p.level, "line_spacing": p.line_spacing,
                      "space_before": p.space_before.pt if p.space_before else None,
                      "space_after": p.space_after.pt if p.space_after else None})
    return {"paragraphs": paras, "word_wrap": tf.word_wrap,
            "vertical_anchor": str(tf.vertical_anchor) if tf.vertical_anchor else None,
            "margins": [tf.margin_left, tf.margin_right, tf.margin_top, tf.margin_bottom]}


def _shape(sh, media_dir, out):
    d = {"name": sh.shape_id and sh.name, "left": sh.left, "top": sh.top, "width": sh.width, "height": sh.height,
         "rotation": getattr(sh, "rotation", 0.0)}
    st = sh.shape_type
    try:
        if st == MSO_GROUP:
            d["kind"] = "group"
            d["shapes"] = [_shape(s2, media_dir, out) for s2 in sh.shapes]
            return d
        if st == MSO_PICTURE:
            blob = sh.image.blob
            h = hashlib.md5(blob).hexdigest()[:16]
            ext = sh.image.ext
            fn = f"{h}.{ext}"
            path = os.path.join(media_dir, fn)
            if not os.path.exists(path):
                with open(path, "wb") as f:
                    f.write(blob)
            d.update(kind="picture", media=fn)
            try:
                d["crop"] = [sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom]
            except Exception:
                pass
            return d
        if st == MSO_TABLE:
            t = sh.table
            d.update(kind="table",
                     rows=[[_text(c.text_frame) for c in row.cells] for row in t.rows],
                     col_widths=[c.width for c in t.columns], row_heights=[r.height for r in t.rows],
                     cell_fills=[[_color(c.fill.fore_color) if str(c.fill.type).find("SOLID") >= 0 else None
                                  for c in row.cells] for row in t.rows],
                     first_row=t.first_row, horz_banding=t.horz_banding)
            return d
        if sh.has_text_frame:
            d.update(kind="shape" if st is not None and "TEXT_BOX" not in str(st) else "textbox",
                     autoshape=str(st), text=_text(sh.text_frame), fill=_fill(sh), line=_line(sh))
            try:
                d["adjustments"] = [a for a in sh.adjustments]
            except Exception:
                pass
            return d
    except Exception as e:  # pragma: no cover
        out.setdefault("warnings", []).append(f"shape {sh.shape_id}: {e}")
    d.update(kind="unsupported", autoshape=str(st), xml=sh._element.xml)
    return d


def main(src, spec_dir):
    os.makedirs(spec_dir, exist_ok=True)
    media = os.path.join(spec_dir, "media")
    os.makedirs(media, exist_ok=True)
    prs = Presentation(src)
    out = {"source": os.path.abspath(src), "slide_width": prs.slide_width, "slide_height": prs.slide_height,
           "slides": []}
    for s in prs.slides:
        bg = None
        try:
            if str(s.background.fill.type).find("SOLID") >= 0:
                bg = _color(s.background.fill.fore_color)
        except Exception:
            pass
        out["slides"].append({
            "background": bg,
            "shapes": [_shape(sh, media, out) for sh in s.shapes],
            "notes": s.notes_slide.notes_text_frame.text if s.has_notes_slide else "",
        })
    path = os.path.join(spec_dir, "deck_spec.json")
    with open(path, "w") as f:
        json.dump(out, f, indent=1)
    kinds = {}
    for sl in out["slides"]:
        def walk(shs):
            for sh in shs:
                kinds[sh["kind"]] = kinds.get(sh["kind"], 0) + 1
                if sh["kind"] == "group":
                    walk(sh["shapes"])
        walk(sl["shapes"])
    print(f"[import] {len(out['slides'])} slides -> {path}")
    print(f"[import] shapes by kind: {kinds}")
    print(f"[import] media files: {len(os.listdir(media))}")
    if out.get("warnings"):
        print(f"[import] warnings: {len(out['warnings'])}")
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("spec_dir")
    a = ap.parse_args()
    main(a.src, a.spec_dir)
